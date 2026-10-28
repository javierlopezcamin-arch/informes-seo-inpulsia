#!/usr/bin/env python3
"""
Generador de presentaciones PPTX — high-end design Inpulsia
"""

import os
import urllib.request
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.oxml.xmlchemy import OxmlElement
from pptx.enum.text import PP_ALIGN
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.dml.color import RGBColor
from pptx.chart.data import CategoryChartData
from utils import MESES, get_output_dir

try:
    import charts as gcharts
except ImportError:
    gcharts = None

# ── Paleta Inpulsia ──────────────────────────────────────────
NAVY  = RGBColor(0x1C, 0x1F, 0x37)
NAVY2 = RGBColor(0x23, 0x27, 0x3F)
NAVY3 = RGBColor(0x14, 0x17, 0x2B)
LIME  = RGBColor(0xC6, 0xFF, 0x2A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
OFFWH = RGBColor(0xF5, 0xF6, 0xFA)
DARK  = RGBColor(0x1A, 0x1C, 0x2E)
MUTED = RGBColor(0x8A, 0x94, 0xB0)
GREEN = RGBColor(0x00, 0xE5, 0xA0)
RED   = RGBColor(0xFF, 0x4D, 0x6A)
WMARK = RGBColor(0xE8, 0xEA, 0xF2)

MESES_UPPER = {k: v.upper() for k, v in MESES.items()}

_CACHE_DIR = os.path.expanduser("~/.cache/seo-reports/")
_LOGO_URL  = "https://inpulsia.es/wp-content/uploads/2025/12/inpulsia-logo-2.png"
_LOGO_PATH = os.path.join(_CACHE_DIR, "inpulsia-logo.png")


# ── Helpers ──────────────────────────────────────────────────

def _fetch(url: str, path: str):
    if os.path.exists(path):
        return path
    try:
        os.makedirs(_CACHE_DIR, exist_ok=True)
        req = urllib.request.Request(
            url,
            headers={'User-Agent': 'Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36'}
        )
        with urllib.request.urlopen(req, timeout=20) as response:
            with open(path, 'wb') as f:
                f.write(response.read())
        return path
    except Exception as e:
        print(f"  ⚠️  Imagen no descargada ({url[:80]}): {type(e).__name__} — la diapositiva sale sin ella")
        return None


def _get_logo():
    return _fetch(_LOGO_URL, _LOGO_PATH)


def _svg_to_png(svg_path: str):
    """Rasteriza un SVG a PNG (PIL/python-pptx no entienden SVG).

    Usa cairosvg si está disponible; si no, cae a `qlmanage` (Quick Look,
    nativo de macOS). Recorta los bordes transparentes para conservar el
    aspect ratio real del logo.
    """
    png_path = svg_path.rsplit(".", 1)[0] + ".png"
    if os.path.exists(png_path):
        return png_path
    try:
        import cairosvg
        cairosvg.svg2png(url=svg_path, write_to=png_path, output_width=800)
    except Exception:
        try:
            import subprocess
            out_dir = os.path.dirname(png_path) or "."
            subprocess.run(
                ["qlmanage", "-t", "-s", "800", "-o", out_dir, svg_path],
                check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
            )
            ql_out = os.path.join(out_dir, os.path.basename(svg_path) + ".png")
            if os.path.exists(ql_out):
                os.replace(ql_out, png_path)
            else:
                print(f"  ⚠️  Logo SVG no convertido ({os.path.basename(svg_path)}) — sale sin logo")
                return None
        except Exception as e:
            print(f"  ⚠️  Logo SVG no convertido ({os.path.basename(svg_path)}): {type(e).__name__} — sale sin logo")
            return None
    # recorta bordes transparentes (qlmanage rellena a un cuadrado)
    try:
        from PIL import Image
        with Image.open(png_path) as im:
            im = im.convert("RGBA")
            bbox = im.getbbox()
            if bbox:
                im.crop(bbox).save(png_path)
    except Exception:
        pass
    return png_path


def _get_client_logo(client_name: str):
    key = f"LOGO_URL_{client_name.upper().replace(' ', '_').replace('-', '_')}"
    url = os.getenv(key, "").strip()
    if not url:
        return None
    ext = url.rsplit(".", 1)[-1].split("?")[0].lower() or "png"
    path = os.path.join(_CACHE_DIR, f"logo-{client_name.lower().replace(' ', '_')}.{ext}")
    path = _fetch(url, path)
    if not path:
        return None
    # SVG por extensión o por contenido → rasterizar a PNG
    is_svg = ext == "svg"
    if not is_svg:
        try:
            with open(path, "rb") as f:
                head = f.read(512).lstrip()
            is_svg = head[:5].lower() == b"<?xml" and b"<svg" in head.lower() \
                or head[:4].lower() == b"<svg"
        except Exception:
            is_svg = False
    if is_svg:
        return _svg_to_png(path)
    return path


def _delta_color(delta: str) -> RGBColor:
    if delta.startswith("▲"):
        return GREEN
    if delta.startswith("▼"):
        return RED
    return MUTED


def _section_watermark(slide, number: str):
    tb = slide.shapes.add_textbox(Inches(6.2), Inches(2.8), Inches(3.9), Inches(2.9))
    p = tb.text_frame.paragraphs[0]
    p.text = number
    p.font.size = Pt(130)
    p.font.bold = True
    p.font.color.rgb = WMARK
    p.alignment = PP_ALIGN.RIGHT


def _add_top_bar(slide, right_text: str = ""):
    bar = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(10), Inches(0.58))
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()

    acc = slide.shapes.add_shape(1, Inches(0), Inches(0.58), Inches(10), Inches(0.035))
    acc.fill.solid()
    acc.fill.fore_color.rgb = LIME
    acc.line.fill.background()

    logo = _get_logo()
    if logo:
        slide.shapes.add_picture(logo, Inches(0.2), Inches(0.09), height=Inches(0.4))

    if right_text:
        tb = slide.shapes.add_textbox(Inches(4.8), Inches(0.07), Inches(5.0), Inches(0.44))
        p = tb.text_frame.paragraphs[0]
        p.text = right_text.upper()
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = LIME
        p.alignment = PP_ALIGN.RIGHT


def _slide_title(slide, text: str):
    acc = slide.shapes.add_shape(1, Inches(0.42), Inches(0.76), Inches(0.055), Inches(0.5))
    acc.fill.solid()
    acc.fill.fore_color.rgb = LIME
    acc.line.fill.background()

    tb = slide.shapes.add_textbox(Inches(0.6), Inches(0.76), Inches(9.0), Inches(0.55))
    p = tb.text_frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = DARK


# ── Diapositivas ─────────────────────────────────────────────

def add_cover_slide(prs, client_name: str, month: int, year: int, client_logo=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY

    # Banda lima vertical izquierda
    band = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(0.22), Inches(5.625))
    band.fill.solid()
    band.fill.fore_color.rgb = LIME
    band.line.fill.background()

    # Logo Inpulsia — centrado arriba
    logo = _get_logo()
    if logo:
        pic = slide.shapes.add_picture(logo, Inches(0), Inches(0.22), height=Inches(0.5))
        pic.left = int((prs.slide_width - pic.width) / 2)

    # "INFORME SEO"
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(1.05), Inches(9.2), Inches(1.3))
    p = tb.text_frame.paragraphs[0]
    p.text = "INFORME SEO"
    p.font.size = Pt(60)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER

    # Filete lima separador
    sep = slide.shapes.add_shape(1, Inches(3.8), Inches(2.45), Inches(2.4), Inches(0.055))
    sep.fill.solid()
    sep.fill.fore_color.rgb = LIME
    sep.line.fill.background()

    # Logo cliente en círculo blanco centrado (si existe)
    y_name = 2.62
    if client_logo:
        from pptx.enum.shapes import MSO_SHAPE
        circle_d = 1.2
        circle_x = (10 - circle_d) / 2
        circle = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(circle_x), Inches(y_name),
            Inches(circle_d), Inches(circle_d))
        circle.fill.solid()
        circle.fill.fore_color.rgb = WHITE
        circle.line.fill.background()
        logo_h = 0.7
        pic = slide.shapes.add_picture(client_logo, Inches(0), Inches(0), height=Inches(logo_h))
        pic.left = int((prs.slide_width - pic.width) / 2)
        pic.top = int(Inches(y_name) + (Inches(circle_d) - pic.height) / 2)
        y_name += circle_d + 0.15

    # Nombre cliente
    tb2 = slide.shapes.add_textbox(Inches(0.5), Inches(y_name), Inches(9.2), Inches(0.7))
    p2 = tb2.text_frame.paragraphs[0]
    p2.text = client_name.upper()
    p2.font.size = Pt(28)
    p2.font.bold = True
    p2.font.color.rgb = LIME
    p2.alignment = PP_ALIGN.CENTER

    # Mes / año
    tb3 = slide.shapes.add_textbox(Inches(0.5), Inches(y_name + 0.72), Inches(9.2), Inches(0.5))
    p3 = tb3.text_frame.paragraphs[0]
    p3.text = f"{MESES_UPPER.get(month, 'MES')}  {year}"
    p3.font.size = Pt(18)
    p3.font.color.rgb = MUTED
    p3.alignment = PP_ALIGN.CENTER

    # Pie NAVY2
    foot = slide.shapes.add_shape(1, Inches(0), Inches(5.05), Inches(10), Inches(0.575))
    foot.fill.solid()
    foot.fill.fore_color.rgb = NAVY2
    foot.line.fill.background()
    tb_f = slide.shapes.add_textbox(Inches(0.5), Inches(5.1), Inches(9.2), Inches(0.4))
    pf = tb_f.text_frame.paragraphs[0]
    pf.text = "INPULSIA · CONSULTORÍA SEO"
    pf.font.size = Pt(9)
    pf.font.bold = True
    pf.font.color.rgb = MUTED
    pf.alignment = PP_ALIGN.CENTER


def add_kpi_slide(prs, title: str, narrative: str, kpis: list, section: str, section_num: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = OFFWH

    _add_top_bar(slide, right_text=section)
    _section_watermark(slide, section_num)
    _slide_title(slide, title)

    # Narrativa
    tb_n = slide.shapes.add_textbox(Inches(0.45), Inches(1.35), Inches(9.1), Inches(0.6))
    tb_n.text_frame.word_wrap = True
    p_n = tb_n.text_frame.paragraphs[0]
    p_n.text = narrative
    p_n.font.size = Pt(11)
    p_n.font.color.rgb = MUTED

    # KPI cards
    n = len(kpis)
    margin = 0.4
    gap = 0.16
    card_w = (10 - 2 * margin - gap * (n - 1)) / n
    card_h = 2.3
    card_y = 2.1

    for i, kpi in enumerate(kpis):
        x = margin + i * (card_w + gap)
        delta = kpi.get("delta", "")
        dcol = _delta_color(delta)

        # Card navy
        card = slide.shapes.add_shape(1, Inches(x), Inches(card_y), Inches(card_w), Inches(card_h))
        card.fill.solid()
        card.fill.fore_color.rgb = NAVY
        card.line.fill.background()

        # Accent top (color según delta)
        top = slide.shapes.add_shape(1, Inches(x), Inches(card_y), Inches(card_w), Inches(0.06))
        top.fill.solid()
        top.fill.fore_color.rgb = dcol
        top.line.fill.background()

        # Label
        tb_l = slide.shapes.add_textbox(Inches(x + 0.17), Inches(card_y + 0.17), Inches(card_w - 0.22), Inches(0.3))
        p_l = tb_l.text_frame.paragraphs[0]
        p_l.text = kpi["label"].upper()
        p_l.font.size = Pt(8)
        p_l.font.bold = True
        p_l.font.color.rgb = MUTED

        # Valor
        tb_v = slide.shapes.add_textbox(Inches(x + 0.14), Inches(card_y + 0.5), Inches(card_w - 0.18), Inches(0.95))
        p_v = tb_v.text_frame.paragraphs[0]
        p_v.text = str(kpi["value"])
        p_v.font.size = Pt(32)
        p_v.font.bold = True
        p_v.font.color.rgb = WHITE

        # Delta
        if delta:
            tb_d = slide.shapes.add_textbox(Inches(x + 0.14), Inches(card_y + 1.55), Inches(card_w - 0.18), Inches(0.42))
            p_d = tb_d.text_frame.paragraphs[0]
            p_d.text = f"{delta}  vs mes anterior"
            p_d.font.size = Pt(9)
            p_d.font.color.rgb = dcol


def add_text_slide(prs, title: str, content: str, section: str = "", section_num: str = ""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = OFFWH

    _add_top_bar(slide, right_text=section)
    if section_num:
        _section_watermark(slide, section_num)
    _slide_title(slide, title)

    box = slide.shapes.add_textbox(Inches(0.45), Inches(1.45), Inches(9.1), Inches(3.9))
    frame = box.text_frame
    frame.word_wrap = True
    frame.clear()
    for line in content.split('\n'):
        if line.strip():
            p = frame.add_paragraph()
            p.text = line.strip()
            p.font.size = Pt(12.5)
            p.font.color.rgb = DARK
            p.space_before = Pt(4)
            p.space_after = Pt(3)


def add_chart_slide(prs, title: str, data_dict: dict, section: str = "", section_num: str = ""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = OFFWH

    _add_top_bar(slide, right_text=section)
    if section_num:
        _section_watermark(slide, section_num)
    _slide_title(slide, title)

    chart_data = CategoryChartData()
    chart_data.categories = list(data_dict.keys())
    chart_data.add_series("", tuple(data_dict.values()))

    cs = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(0.45), Inches(1.38), Inches(9.1), Inches(3.95),
        chart_data
    )
    chart = cs.chart
    chart.has_legend = False
    chart.chart_title.text_frame.clear()

    # Fondos (utilizando manipulación XML ya que python-pptx no expone format para chart/plot area)
    color_hex = "%02X%02X%02X" % (OFFWH[0], OFFWH[1], OFFWH[2])
    
    # 1. Chart area background
    spPr = OxmlElement("c:spPr")
    solidFill = OxmlElement("a:solidFill")
    srgbClr = OxmlElement("a:srgbClr")
    srgbClr.set("val", color_hex)
    solidFill.append(srgbClr)
    spPr.append(solidFill)
    
    ln = OxmlElement("a:ln")
    noFill = OxmlElement("a:noFill")
    ln.append(noFill)
    spPr.append(ln)
    
    chart.element.append(spPr)

    # 2. Plot area background
    plot_spPr = OxmlElement("c:spPr")
    plot_solidFill = OxmlElement("a:solidFill")
    plot_srgbClr = OxmlElement("a:srgbClr")
    plot_srgbClr.set("val", color_hex)
    plot_solidFill.append(plot_srgbClr)
    plot_spPr.append(plot_solidFill)
    
    plot_ln = OxmlElement("a:ln")
    plot_noFill = OxmlElement("a:noFill")
    plot_ln.append(plot_noFill)
    plot_spPr.append(plot_ln)
    
    chart._element.chart.plotArea.append(plot_spPr)

    # Barras lime
    series = chart.plots[0].series[0]
    series.format.fill.solid()
    series.format.fill.fore_color.rgb = LIME

    # Data labels
    series.data_labels.show_value = True
    series.data_labels.font.size = Pt(10)
    series.data_labels.font.bold = True
    series.data_labels.font.color.rgb = DARK
    series.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END


def add_contact_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY

    # Banda lima izquierda (coherencia con portada)
    band = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(0.22), Inches(5.625))
    band.fill.solid()
    band.fill.fore_color.rgb = LIME
    band.line.fill.background()

    # Rombo decorativo esquina superior derecha
    deco = slide.shapes.add_shape(1, Inches(8.3), Inches(-0.6), Inches(2.2), Inches(2.2))
    deco.fill.solid()
    deco.fill.fore_color.rgb = RGBColor(0xC6, 0xFF, 0x2A)
    deco.line.fill.background()
    deco.rotation = 45

    # Logo Inpulsia centrado
    logo = _get_logo()
    if logo:
        pic = slide.shapes.add_picture(logo, Inches(0), Inches(0.6), height=Inches(0.72))
        pic.left = int((prs.slide_width - pic.width) / 2)

    # Separador lima
    sep = slide.shapes.add_shape(1, Inches(3.5), Inches(1.52), Inches(3.0), Inches(0.04))
    sep.fill.solid()
    sep.fill.fore_color.rgb = LIME
    sep.line.fill.background()

    # "Contáctanos"
    tb_c = slide.shapes.add_textbox(Inches(0), Inches(1.65), Inches(10), Inches(0.58))
    p_c = tb_c.text_frame.paragraphs[0]
    p_c.text = "Contáctanos"
    p_c.font.size = Pt(28)
    p_c.font.bold = True
    p_c.font.color.rgb = LIME
    p_c.alignment = PP_ALIGN.CENTER

    # Datos de contacto — texto centrado, sin iconos
    contact_lines = [
        ("hola@inpulsia.es", True),
        ("(+34) 942 135 917  ·  (+34) 609 340 441", False),
        ("", False),
        ("Avda. Eduardo García del Río, 39", False),
        ("39011 · Santander · Cantabria", False),
        ("", False),
        ("www.inpulsia.es", True),
    ]
    y = 2.42
    for text, is_accent in contact_lines:
        if not text:
            y += 0.18
            continue
        tb = slide.shapes.add_textbox(Inches(0), Inches(y), Inches(10), Inches(0.34))
        p = tb.text_frame.paragraphs[0]
        p.text = text
        p.font.size = Pt(13)
        p.font.color.rgb = LIME if is_accent else WHITE
        p.alignment = PP_ALIGN.CENTER
        y += 0.38

    # Pie
    foot = slide.shapes.add_shape(1, Inches(0), Inches(5.05), Inches(10), Inches(0.575))
    foot.fill.solid()
    foot.fill.fore_color.rgb = NAVY2
    foot.line.fill.background()
    tb_f = slide.shapes.add_textbox(Inches(0), Inches(5.1), Inches(10), Inches(0.4))
    pf = tb_f.text_frame.paragraphs[0]
    pf.text = "INPULSIA · CONSULTORÍA SEO"
    pf.font.size = Pt(9)
    pf.font.bold = True
    pf.font.color.rgb = MUTED
    pf.alignment = PP_ALIGN.CENTER


def add_channels_slide(prs, title: str, channels_data: list, narrative: str, section: str = "", section_num: str = ""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = OFFWH

    _add_top_bar(slide, right_text=section)
    if section_num:
        _section_watermark(slide, section_num)
    _slide_title(slide, title)

    # Narrativa del canal
    tb_n = slide.shapes.add_textbox(Inches(0.45), Inches(1.35), Inches(9.1), Inches(0.6))
    tb_n.text_frame.word_wrap = True
    p_n = tb_n.text_frame.paragraphs[0]
    p_n.text = narrative
    p_n.font.size = Pt(11)
    p_n.font.color.rgb = MUTED

    # Preparar datos (top 5 ordenados de mayor a menor, reversados para orden correcto en eje Y)
    top_channels = sorted(channels_data, key=lambda x: x['sessions'], reverse=True)[:5]
    top_channels.reverse()
    
    chart_data = CategoryChartData()
    chart_data.categories = [c['channel'] for c in top_channels]
    chart_data.add_series("", tuple(c['sessions'] for c in top_channels))

    cs = slide.shapes.add_chart(
        XL_CHART_TYPE.BAR_CLUSTERED,
        Inches(0.45), Inches(2.0), Inches(9.1), Inches(3.2),
        chart_data
    )
    chart = cs.chart
    chart.has_legend = False
    chart.chart_title.text_frame.clear()

    # Fondos
    color_hex = "%02X%02X%02X" % (OFFWH[0], OFFWH[1], OFFWH[2])
    
    spPr = OxmlElement("c:spPr")
    solidFill = OxmlElement("a:solidFill")
    srgbClr = OxmlElement("a:srgbClr")
    srgbClr.set("val", color_hex)
    solidFill.append(srgbClr)
    spPr.append(solidFill)
    ln = OxmlElement("a:ln")
    noFill = OxmlElement("a:noFill")
    ln.append(noFill)
    spPr.append(ln)
    chart.element.append(spPr)

    plot_spPr = OxmlElement("c:spPr")
    plot_solidFill = OxmlElement("a:solidFill")
    plot_srgbClr = OxmlElement("a:srgbClr")
    plot_srgbClr.set("val", color_hex)
    plot_solidFill.append(plot_srgbClr)
    plot_spPr.append(plot_solidFill)
    plot_ln = OxmlElement("a:ln")
    plot_noFill = OxmlElement("a:noFill")
    plot_ln.append(plot_noFill)
    plot_spPr.append(plot_ln)
    chart._element.chart.plotArea.append(plot_spPr)

    # Barras lime
    series = chart.plots[0].series[0]
    series.format.fill.solid()
    series.format.fill.fore_color.rgb = LIME

    # Data labels
    series.data_labels.show_value = True
    series.data_labels.font.size = Pt(9.5)
    series.data_labels.font.bold = True
    series.data_labels.font.color.rgb = DARK
    series.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END


def add_devices_slide(prs, title: str, devices_data: dict, narrative: str, section: str = "", section_num: str = ""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = OFFWH

    _add_top_bar(slide, right_text=section)
    if section_num:
        _section_watermark(slide, section_num)
    _slide_title(slide, title)

    # Narrativa del dispositivo
    tb_n = slide.shapes.add_textbox(Inches(0.45), Inches(1.35), Inches(9.1), Inches(0.6))
    tb_n.text_frame.word_wrap = True
    p_n = tb_n.text_frame.paragraphs[0]
    p_n.text = narrative
    p_n.font.size = Pt(11)
    p_n.font.color.rgb = MUTED

    # Preparar datos
    chart_data = CategoryChartData()
    chart_data.categories = list(devices_data.keys())
    chart_data.add_series("", tuple(devices_data.values()))

    # Gráfico Doughnut centrado
    cs = slide.shapes.add_chart(
        XL_CHART_TYPE.DOUGHNUT,
        Inches(2.5), Inches(2.0), Inches(5.0), Inches(3.2),
        chart_data
    )
    chart = cs.chart
    chart.has_legend = True
    chart.legend.font.size = Pt(9.5)
    chart.legend.font.color.rgb = DARK
    chart.chart_title.text_frame.clear()

    # Fondos
    color_hex = "%02X%02X%02X" % (OFFWH[0], OFFWH[1], OFFWH[2])
    
    spPr = OxmlElement("c:spPr")
    solidFill = OxmlElement("a:solidFill")
    srgbClr = OxmlElement("a:srgbClr")
    srgbClr.set("val", color_hex)
    solidFill.append(srgbClr)
    spPr.append(solidFill)
    ln = OxmlElement("a:ln")
    noFill = OxmlElement("a:noFill")
    ln.append(noFill)
    spPr.append(ln)
    chart.element.append(spPr)

    # Formatear sectores con colores premium de Inpulsia
    colors = [LIME, NAVY, MUTED, RED, GREEN]
    points = chart.plots[0].series[0].points
    for idx, pt in enumerate(points):
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = colors[idx % len(colors)]

    # Mostrar porcentajes
    chart.plots[0].has_data_labels = True
    data_labels = chart.plots[0].data_labels
    data_labels.show_percentage = True
    data_labels.show_value = False
    data_labels.font.size = Pt(9)
    data_labels.font.bold = True
    data_labels.font.color.rgb = WHITE


def add_panel_screenshot_slides(prs, client_name: str):
    """Escanea la carpeta de capturas de pantalla e inserta diapositivas con las capturas encontradas"""
    import glob
    base = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    client_dir = os.path.join(base, "capturas", client_name.upper().replace(" ", "_"))
    legacy_dir = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))), "_archivo", "capturas")
    capturas_dir = client_dir if os.path.isdir(client_dir) else legacy_dir
    if not os.path.exists(capturas_dir):
        os.makedirs(capturas_dir, exist_ok=True)
        readme_path = os.path.join(capturas_dir, "LEEME.txt")
        if not os.path.exists(readme_path):
            with open(readme_path, "w", encoding="utf-8") as f:
                f.write(
                    "SISTEMA DE CAPTURAS DE PANTALLA\n"
                    "===============================\n\n"
                    "Coloca aquí las capturas de pantalla de los paneles oficiales para que se inserten automáticamente en el informe.\n"
                    "El script buscará cualquier imagen (PNG, JPG, JPEG) y creará una diapositiva enmarcada para cada una.\n\n"
                    "Ejemplos recomendados:\n"
                    "- analytics_panel.png\n"
                    "- search_console_panel.png\n"
                )
    
    # Buscar imágenes
    extensions = ("*.png", "*.jpg", "*.jpeg")
    images = []
    for ext in extensions:
        images.extend(glob.glob(os.path.join(capturas_dir, ext)))
        images.extend(glob.glob(os.path.join(capturas_dir, ext.upper())))
    
    if not images:
        return
        
    print(f"  📸 Encontradas {len(images)} capturas de pantalla para insertar")
    
    for img_path in sorted(images):
        filename = os.path.basename(img_path)
        if filename.lower() == "leeme.txt":
            continue
            
        name_clean = os.path.splitext(filename)[0].replace("_", " ").replace("-", " ").title()
        
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = OFFWH
        
        _add_top_bar(slide, right_text="Panel Oficial")
        _slide_title(slide, f"Panel: {name_clean}")
        
        # Imagen real respetando proporción, centrada con marco navy
        from PIL import Image as _PILImage
        try:
            with _PILImage.open(img_path) as im:
                ratio = im.width / im.height
        except Exception:
            ratio = 16 / 9
        max_w, max_h = 8.6, 3.85
        w = min(max_w, max_h * ratio)
        h = w / ratio
        x = (10 - w) / 2
        y = 1.42 + (max_h - h) / 2

        border = slide.shapes.add_shape(1, Inches(x - 0.06), Inches(y - 0.06),
                                        Inches(w + 0.12), Inches(h + 0.12))
        border.fill.solid()
        border.fill.fore_color.rgb = NAVY
        border.line.fill.background()

        slide.shapes.add_picture(img_path, Inches(x), Inches(y), width=Inches(w), height=Inches(h))


def add_index_slide(prs, sections: list):
    """Índice del informe — fondo navy, secciones numeradas"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY

    band = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(0.22), Inches(5.625))
    band.fill.solid()
    band.fill.fore_color.rgb = LIME
    band.line.fill.background()

    logo = _get_logo()
    if logo:
        slide.shapes.add_picture(logo, Inches(0.42), Inches(0.22), height=Inches(0.45))

    tb = slide.shapes.add_textbox(Inches(0.6), Inches(0.95), Inches(8.8), Inches(0.7))
    p = tb.text_frame.paragraphs[0]
    p.text = "Contenido"
    p.font.size = Pt(34)
    p.font.bold = True
    p.font.color.rgb = WHITE

    y = 1.85
    for num, name in sections:
        tb_n = slide.shapes.add_textbox(Inches(0.85), Inches(y), Inches(0.7), Inches(0.45))
        p_n = tb_n.text_frame.paragraphs[0]
        p_n.text = num
        p_n.font.size = Pt(20)
        p_n.font.bold = True
        p_n.font.color.rgb = LIME

        tb_s = slide.shapes.add_textbox(Inches(1.65), Inches(y + 0.04), Inches(7.6), Inches(0.45))
        p_s = tb_s.text_frame.paragraphs[0]
        p_s.text = name
        p_s.font.size = Pt(15)
        p_s.font.color.rgb = WHITE
        y += 0.55


def add_image_chart_slide(prs, title: str, img_path: str, narrative: str = "",
                          section: str = "", section_num: str = ""):
    """Slide con gráfico matplotlib (PNG) a página completa"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = OFFWH

    _add_top_bar(slide, right_text=section)
    if section_num:
        _section_watermark(slide, section_num)
    _slide_title(slide, title)

    y_img = 1.4
    if narrative:
        tb_n = slide.shapes.add_textbox(Inches(0.45), Inches(1.32), Inches(9.1), Inches(0.5))
        tb_n.text_frame.word_wrap = True
        p_n = tb_n.text_frame.paragraphs[0]
        p_n.text = narrative
        p_n.font.size = Pt(11)
        p_n.font.color.rgb = MUTED
        y_img = 1.85

    from PIL import Image as _PILImage
    try:
        with _PILImage.open(img_path) as im:
            ratio = im.width / im.height
    except Exception:
        ratio = 2.6
    max_w, max_h = 9.1, 5.45 - y_img - 0.15
    w = min(max_w, max_h * ratio)
    h = w / ratio
    x = (10 - w) / 2
    slide.shapes.add_picture(img_path, Inches(x), Inches(y_img), width=Inches(w), height=Inches(h))


def add_double_image_slide(prs, title: str, img_left: str, label_left: str,
                           img_right: str, label_right: str,
                           section: str = "", section_num: str = ""):
    """Slide con dos gráficos lado a lado (dispositivos + países)"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = OFFWH

    _add_top_bar(slide, right_text=section)
    if section_num:
        _section_watermark(slide, section_num)
    _slide_title(slide, title)

    for img, label, x in ((img_left, label_left, 0.45), (img_right, label_right, 5.15)):
        if not img:
            continue
        tb = slide.shapes.add_textbox(Inches(x), Inches(1.42), Inches(4.4), Inches(0.32))
        p = tb.text_frame.paragraphs[0]
        p.text = label.upper()
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = MUTED
        slide.shapes.add_picture(img, Inches(x), Inches(1.8), width=Inches(4.4))


def add_table_slide(prs, title: str, headers: list, rows: list, narrative: str = "",
                    section: str = "", section_num: str = "", col_widths: list = None):
    """Tabla premium: cabecera navy, filas alternas, sin bordes pesados"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = OFFWH

    _add_top_bar(slide, right_text=section)
    if section_num:
        _section_watermark(slide, section_num)
    _slide_title(slide, title)

    y_table = 1.42
    if narrative:
        tb_n = slide.shapes.add_textbox(Inches(0.45), Inches(1.34), Inches(9.1), Inches(0.45))
        tb_n.text_frame.word_wrap = True
        p_n = tb_n.text_frame.paragraphs[0]
        p_n.text = narrative
        p_n.font.size = Pt(11)
        p_n.font.color.rgb = MUTED
        y_table = 1.88

    n_rows = len(rows) + 1
    n_cols = len(headers)
    table_h = min(3.55, 0.34 * n_rows)
    gframe = slide.shapes.add_table(n_rows, n_cols, Inches(0.45), Inches(y_table),
                                    Inches(9.1), Inches(table_h))
    table = gframe.table

    if col_widths:
        total = sum(col_widths)
        for i, cw in enumerate(col_widths):
            table.columns[i].width = Inches(9.1 * cw / total)

    # Cabecera
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        cell.margin_top = cell.margin_bottom = Inches(0.03)
        p = cell.text_frame.paragraphs[0]
        p.text = h.upper()
        p.font.size = Pt(9)
        p.font.bold = True
        p.font.color.rgb = LIME

    # Filas
    for i, row in enumerate(rows, 1):
        for j, val in enumerate(row):
            cell = table.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if i % 2 else RGBColor(0xEE, 0xF0, 0xF6)
            cell.margin_top = cell.margin_bottom = Inches(0.02)
            p = cell.text_frame.paragraphs[0]
            p.text = str(val)
            p.font.size = Pt(9.5)
            p.font.color.rgb = DARK
            if j > 0:
                p.alignment = PP_ALIGN.RIGHT


def add_pagespeed_slide(prs, ps_data: dict, section: str = "", section_num: str = ""):
    """Slide de SEO técnico con semáforo Performance / SEO / Accessibility."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = OFFWH

    _add_top_bar(slide, right_text=section)
    if section_num:
        _section_watermark(slide, section_num)
    _slide_title(slide, "SEO Técnico — PageSpeed Insights")

    tb_n = slide.shapes.add_textbox(Inches(0.45), Inches(1.35), Inches(9.1), Inches(0.5))
    tb_n.text_frame.word_wrap = True
    p_n = tb_n.text_frame.paragraphs[0]
    p_n.text = "Puntuaciones de rendimiento, SEO técnico y accesibilidad del sitio web según Google PageSpeed Insights."
    p_n.font.size = Pt(11)
    p_n.font.color.rgb = MUTED

    def _score_color(score):
        if score >= 90: return GREEN
        if score >= 50: return RGBColor(0xFB, 0xBC, 0x04)
        return RED

    scores = [
        ("PERFORMANCE", ps_data.get("performance", 0)),
        ("SEO", ps_data.get("seo", 0)),
        ("ACCESIBILIDAD", ps_data.get("accessibility", 0)),
    ]

    card_w = 2.7
    gap = 0.3
    start_x = (10 - len(scores) * card_w - (len(scores) - 1) * gap) / 2

    for i, (label, score) in enumerate(scores):
        x = start_x + i * (card_w + gap)
        card = slide.shapes.add_shape(1, Inches(x), Inches(2.2), Inches(card_w), Inches(1.8))
        card.fill.solid()
        card.fill.fore_color.rgb = NAVY
        card.line.fill.background()

        top = slide.shapes.add_shape(1, Inches(x), Inches(2.2), Inches(card_w), Inches(0.06))
        top.fill.solid()
        top.fill.fore_color.rgb = _score_color(score)
        top.line.fill.background()

        tb_l = slide.shapes.add_textbox(Inches(x + 0.15), Inches(2.35), Inches(card_w - 0.3), Inches(0.3))
        p_l = tb_l.text_frame.paragraphs[0]
        p_l.text = label
        p_l.font.size = Pt(9)
        p_l.font.bold = True
        p_l.font.color.rgb = MUTED

        tb_v = slide.shapes.add_textbox(Inches(x + 0.15), Inches(2.7), Inches(card_w - 0.3), Inches(0.85))
        p_v = tb_v.text_frame.paragraphs[0]
        p_v.text = str(score)
        p_v.font.size = Pt(42)
        p_v.font.bold = True
        p_v.font.color.rgb = _score_color(score)
        p_v.alignment = PP_ALIGN.CENTER

    cwv = [
        ("FCP", ps_data.get("fcp", "—")),
        ("LCP", ps_data.get("lcp", "—")),
        ("CLS", ps_data.get("cls", "—")),
        ("TBT", ps_data.get("tbt", "—")),
    ]
    tb_cwv = slide.shapes.add_textbox(Inches(0.45), Inches(4.3), Inches(9.1), Inches(0.5))
    tb_cwv.text_frame.word_wrap = True
    p_cwv = tb_cwv.text_frame.paragraphs[0]
    p_cwv.text = "Core Web Vitals:  " + "  ·  ".join(f"{k}: {v}" for k, v in cwv)
    p_cwv.font.size = Pt(10)
    p_cwv.font.color.rgb = MUTED


def add_executive_slide(prs, summary: str, month: int, year: int):
    """Resumen ejecutivo — slide 2, fondo navy, 3 frases con acento lima."""
    if not summary:
        return
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY

    band = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(0.22), Inches(5.625))
    band.fill.solid()
    band.fill.fore_color.rgb = LIME
    band.line.fill.background()

    logo = _get_logo()
    if logo:
        slide.shapes.add_picture(logo, Inches(0.42), Inches(0.22), height=Inches(0.45))

    tb_t = slide.shapes.add_textbox(Inches(0.6), Inches(0.95), Inches(8.8), Inches(0.55))
    p_t = tb_t.text_frame.paragraphs[0]
    p_t.text = "Resumen Ejecutivo"
    p_t.font.size = Pt(28)
    p_t.font.bold = True
    p_t.font.color.rgb = WHITE

    tb_p = slide.shapes.add_textbox(Inches(0.6), Inches(1.55), Inches(5), Inches(0.3))
    p_p = tb_p.text_frame.paragraphs[0]
    p_p.text = f"{MESES_UPPER.get(month, 'MES')} {year}"
    p_p.font.size = Pt(12)
    p_p.font.color.rgb = MUTED

    lines = [l.strip() for l in summary.split('\n') if l.strip()]
    y = 2.2
    for line in lines[:3]:
        acc = slide.shapes.add_shape(1, Inches(0.6), Inches(y + 0.08), Inches(0.06), Inches(0.22))
        acc.fill.solid()
        acc.fill.fore_color.rgb = LIME
        acc.line.fill.background()

        tb = slide.shapes.add_textbox(Inches(0.85), Inches(y), Inches(8.5), Inches(0.5))
        tb.text_frame.word_wrap = True
        p = tb.text_frame.paragraphs[0]
        p.text = line
        p.font.size = Pt(15)
        p.font.color.rgb = OFFWH
        p.space_after = Pt(6)
        y += 0.65


def _add_insight_box(slide, text: str, y: float = 4.85):
    """Caja de insight IA bajo un gráfico/tabla existente."""
    box = slide.shapes.add_shape(1, Inches(0.45), Inches(y), Inches(9.1), Inches(0.42))
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(0x1E, 0x22, 0x3A)
    box.line.fill.background()

    acc = slide.shapes.add_shape(1, Inches(0.45), Inches(y), Inches(0.05), Inches(0.42))
    acc.fill.solid()
    acc.fill.fore_color.rgb = LIME
    acc.line.fill.background()

    tb = slide.shapes.add_textbox(Inches(0.65), Inches(y + 0.06), Inches(8.8), Inches(0.3))
    tb.text_frame.word_wrap = True
    p = tb.text_frame.paragraphs[0]
    run_ia = p.add_run()
    run_ia.text = "NOTA  "
    run_ia.font.size = Pt(8)
    run_ia.font.bold = True
    run_ia.font.color.rgb = LIME
    run_t = p.add_run()
    run_t.text = text
    run_t.font.size = Pt(9.5)
    run_t.font.italic = True
    run_t.font.color.rgb = OFFWH


def add_keyword_movements_slide(prs, movements: dict, section: str = "", section_num: str = ""):
    gained = movements.get('gained', [])
    lost = movements.get('lost', [])
    if not gained and not lost:
        return
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = OFFWH

    _add_top_bar(slide, right_text=section)
    if section_num:
        _section_watermark(slide, section_num)
    _slide_title(slide, "Keywords — Ganadas y Perdidas")

    tb_n = slide.shapes.add_textbox(Inches(0.45), Inches(1.35), Inches(9.1), Inches(0.4))
    tb_n.text_frame.word_wrap = True
    p_n = tb_n.text_frame.paragraphs[0]
    p_n.text = f"{len(gained)} keywords subieron de posición y {len(lost)} bajaron respecto al mes anterior."
    p_n.font.size = Pt(11)
    p_n.font.color.rgb = MUTED

    def _draw_col(items, x_start, col_w, label, accent_color, arrow):
        tb_h = slide.shapes.add_textbox(Inches(x_start), Inches(1.88), Inches(col_w), Inches(0.32))
        p_h = tb_h.text_frame.paragraphs[0]
        p_h.text = label.upper()
        p_h.font.size = Pt(10)
        p_h.font.bold = True
        p_h.font.color.rgb = accent_color

        for i, kw in enumerate(items[:6]):
            y = 2.28 + i * 0.48
            card = slide.shapes.add_shape(1, Inches(x_start), Inches(y), Inches(col_w), Inches(0.42))
            card.fill.solid()
            card.fill.fore_color.rgb = NAVY
            card.line.fill.background()

            tb_q = slide.shapes.add_textbox(Inches(x_start + 0.14), Inches(y + 0.04), Inches(col_w - 1.6), Inches(0.34))
            p_q = tb_q.text_frame.paragraphs[0]
            p_q.text = kw['query']
            p_q.font.size = Pt(10)
            p_q.font.color.rgb = WHITE

            tb_d = slide.shapes.add_textbox(Inches(x_start + col_w - 1.4), Inches(y + 0.04), Inches(1.26), Inches(0.34))
            p_d = tb_d.text_frame.paragraphs[0]
            delta = abs(kw['delta'])
            p_d.text = f"pos {kw['position']}  {arrow}{delta}"
            p_d.font.size = Pt(9)
            p_d.font.bold = True
            p_d.font.color.rgb = accent_color
            p_d.alignment = PP_ALIGN.RIGHT

    _draw_col(gained, 0.45, 4.4, "Subieron", GREEN, "▲")
    _draw_col(lost, 5.15, 4.4, "Bajaron", RED, "▼")


def add_opportunities_slide(prs, opps_data: dict, section: str = "", section_num: str = ""):
    items = opps_data.get('opportunities', [])
    if not items:
        return
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = OFFWH

    _add_top_bar(slide, right_text=section)
    if section_num:
        _section_watermark(slide, section_num)
    _slide_title(slide, "Oportunidades — Página 2 de Google")

    tb_n = slide.shapes.add_textbox(Inches(0.45), Inches(1.35), Inches(9.1), Inches(0.4))
    tb_n.text_frame.word_wrap = True
    p_n = tb_n.text_frame.paragraphs[0]
    p_n.text = opps_data.get('narrative', '')
    p_n.font.size = Pt(11)
    p_n.font.color.rgb = MUTED

    headers = ["Keyword", "Posición", "Impresiones", "Gap al Top 10"]
    rows = [[it['query'], f"{it['position']}", f"{it['impressions']:,}", f"+{it['gap']} pos."] for it in items]

    n_rows = len(rows) + 1
    n_cols = len(headers)
    table_h = min(3.3, 0.36 * n_rows)
    gframe = slide.shapes.add_table(n_rows, n_cols, Inches(0.45), Inches(1.92),
                                    Inches(9.1), Inches(table_h))
    table = gframe.table
    widths = [4.5, 1.3, 1.8, 1.5]
    total = sum(widths)
    for i, cw in enumerate(widths):
        table.columns[i].width = Inches(9.1 * cw / total)

    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        cell.margin_top = cell.margin_bottom = Inches(0.03)
        p = cell.text_frame.paragraphs[0]
        p.text = h.upper()
        p.font.size = Pt(9)
        p.font.bold = True
        p.font.color.rgb = LIME

    for i, row in enumerate(rows, 1):
        for j, val in enumerate(row):
            cell = table.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if i % 2 else RGBColor(0xEE, 0xF0, 0xF6)
            cell.margin_top = cell.margin_bottom = Inches(0.02)
            p = cell.text_frame.paragraphs[0]
            p.text = str(val)
            p.font.size = Pt(9.5)
            p.font.color.rgb = DARK
            if j == 3:
                p.font.color.rgb = RGBColor(0xFB, 0xBC, 0x04)
                p.font.bold = True
            if j > 0:
                p.alignment = PP_ALIGN.RIGHT


# ── Orquestador ──────────────────────────────────────────────

def convert_to_pdf(pptx_path: str) -> str | None:
    """Convierte PPTX a PDF usando LibreOffice headless. Devuelve ruta del PDF o None."""
    import subprocess, shutil
    soffice = shutil.which("soffice")
    if not soffice:
        print("  ℹ️  PDF: sin LibreOffice, solo se genera el PPTX")
        return None
    out_dir = os.path.dirname(pptx_path)
    try:
        # El primer arranque de LibreOffice en frío puede superar el minuto
        r = subprocess.run([soffice, "--headless", "--convert-to", "pdf", pptx_path, "--outdir", out_dir],
                           capture_output=True, text=True, timeout=150)
        pdf_path = pptx_path.rsplit(".", 1)[0] + ".pdf"
        if os.path.exists(pdf_path):
            print(f"  ✅ PDF generado: {pdf_path}")
            return pdf_path
        err = (r.stderr or r.stdout or "").strip().splitlines()
        print(f"  ⚠️  PDF: LibreOffice no generó el archivo{(' — ' + err[-1]) if err else ''}")
    except subprocess.TimeoutExpired:
        print("  ⚠️  PDF: LibreOffice tardó demasiado (se omite el PDF)")
    except Exception as e:
        print(f"  ⚠️  PDF: error — {e}")
    return None


def add_recommendations_slide(prs, recommendations, anomalies, section, section_num):
    if anomalies:
        _add_recs_slide_anomalies(prs, anomalies, section, section_num)
        _add_recs_slide_recs(prs, recommendations, section, section_num)
    else:
        _add_recs_slide_recs(prs, recommendations, section, section_num)


def _add_recs_slide_anomalies(prs, anomalies, section, section_num):
    """Slide 1/2 — Puntos de atención (hasta 5 items)."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = OFFWH
    _add_top_bar(slide, right_text=section)
    _section_watermark(slide, section_num)
    _slide_title(slide, "Puntos de Atención")

    y = 1.5
    tb = slide.shapes.add_textbox(Inches(0.45), Inches(y), Inches(9.1), Inches(0.3))
    p = tb.text_frame.paragraphs[0]
    p.text = "PUNTOS DE ATENCIÓN"
    p.font.size = Pt(10); p.font.bold = True; p.font.color.rgb = MUTED
    y += 0.4
    for a in anomalies[:5]:
        dot = slide.shapes.add_shape(1, Inches(0.45), Inches(y + 0.05), Inches(0.05), Inches(0.3))
        dot.fill.solid(); dot.fill.fore_color.rgb = RGBColor(0xFB, 0xBC, 0x04); dot.line.fill.background()
        tb = slide.shapes.add_textbox(Inches(0.65), Inches(y), Inches(8.9), Inches(0.4))
        tb.text_frame.word_wrap = True
        p = tb.text_frame.paragraphs[0]
        p.text = a
        p.font.size = Pt(11.5); p.font.color.rgb = DARK
        y += 0.5


def _add_recs_slide_recs(prs, recommendations, section, section_num):
    """Slide 2/2 (o única) — Recomendaciones (hasta 5 items)."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = OFFWH
    _add_top_bar(slide, right_text=section)
    _section_watermark(slide, section_num)
    _slide_title(slide, "Recomendaciones")

    y = 1.5
    tb = slide.shapes.add_textbox(Inches(0.45), Inches(y), Inches(9.1), Inches(0.3))
    p = tb.text_frame.paragraphs[0]
    p.text = "RECOMENDACIONES"
    p.font.size = Pt(10); p.font.bold = True; p.font.color.rgb = MUTED
    y += 0.45
    for i, rec in enumerate(recommendations[:5], 1):
        num = slide.shapes.add_shape(1, Inches(0.45), Inches(y), Inches(0.38), Inches(0.38))
        num.fill.solid(); num.fill.fore_color.rgb = NAVY; num.line.fill.background()
        pn = num.text_frame.paragraphs[0]
        pn.alignment = PP_ALIGN.CENTER
        rn = pn.add_run(); rn.text = str(i); rn.font.size = Pt(13); rn.font.bold = True; rn.font.color.rgb = LIME
        tb = slide.shapes.add_textbox(Inches(1.0), Inches(y), Inches(8.5), Inches(0.5))
        tb.text_frame.word_wrap = True
        p = tb.text_frame.paragraphs[0]
        p.text = rec
        p.font.size = Pt(12); p.font.color.rgb = DARK
        y += 0.62


GLOSSARY = [
    ("Usuarios", "Personas únicas que visitan tu web en el periodo."),
    ("Sesiones", "Visitas a la web; un usuario puede generar varias."),
    ("Páginas vistas", "Número total de páginas cargadas en la web."),
    ("Tasa de rebote", "Porcentaje de visitas que se van sin interactuar."),
    ("Impresiones", "Veces que tu web apareció en los resultados de Google."),
    ("Clics", "Veces que los usuarios pulsaron tu resultado en Google."),
    ("CTR", "Porcentaje de impresiones que se convierten en clic."),
    ("Posición media", "Lugar promedio de tu web en los resultados de búsqueda."),
]


def add_glossary_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = OFFWH

    _add_top_bar(slide, right_text="Glosario")
    _slide_title(slide, "Glosario de Métricas")

    tb_n = slide.shapes.add_textbox(Inches(0.45), Inches(1.35), Inches(9.1), Inches(0.5))
    p_n = tb_n.text_frame.paragraphs[0]
    p_n.text = "Qué significa cada métrica que aparece en este informe."
    p_n.font.size = Pt(11)
    p_n.font.color.rgb = MUTED

    half = (len(GLOSSARY) + 1) // 2
    cols = [(0.45, GLOSSARY[:half]), (5.05, GLOSSARY[half:])]
    for x0, items in cols:
        y = 2.0
        for term, desc in items:
            acc = slide.shapes.add_shape(1, Inches(x0), Inches(y + 0.04), Inches(0.05), Inches(0.62))
            acc.fill.solid()
            acc.fill.fore_color.rgb = LIME
            acc.line.fill.background()

            tb_t = slide.shapes.add_textbox(Inches(x0 + 0.2), Inches(y), Inches(4.2), Inches(0.3))
            p_t = tb_t.text_frame.paragraphs[0]
            p_t.text = term
            p_t.font.size = Pt(12)
            p_t.font.bold = True
            p_t.font.color.rgb = DARK

            tb_d = slide.shapes.add_textbox(Inches(x0 + 0.2), Inches(y + 0.3), Inches(4.2), Inches(0.5))
            tb_d.text_frame.word_wrap = True
            p_d = tb_d.text_frame.paragraphs[0]
            p_d.text = desc
            p_d.font.size = Pt(9.5)
            p_d.font.color.rgb = MUTED
            y += 0.85


def generate_pptx(analysis: dict, client_name: str, sections: set = None, include_notes: bool = True) -> str:
    ALL = {'resumen','trafico','canales','search','keywords','pagespeed','conclusiones'}
    sec = sections if sections else ALL
    month   = analysis['month']
    year    = analysis['year']
    traffic = analysis['traffic']
    sc      = analysis['search_console']
    kw      = analysis['keywords']

    client_logo = _get_client_logo(client_name)

    prs = Presentation()
    prs.slide_width  = Inches(10)
    prs.slide_height = Inches(5.625)

    daily_users   = analysis.get('daily_users', [])
    daily_sc      = analysis.get('daily_trends', {}).get('daily', []) or analysis.get('daily_sc', [])
    top_countries = analysis.get('top_countries', [])
    channels_d    = analysis.get('traffic_channels', {})
    devices_d     = analysis.get('devices', {})
    pagespeed     = analysis.get('pagespeed', {})
    ads           = analysis.get('ads', {})
    mom           = analysis.get('mom', {})
    kw_moves      = analysis.get('keyword_movements', {})
    opps          = analysis.get('opportunities', {})
    insights      = analysis.get('insights', {})
    exec_summary  = analysis.get('executive_summary', '')
    historical    = analysis.get('historical', [])

    # 1. PORTADA
    add_cover_slide(prs, client_name, month, year, client_logo)

    # 1b. RESUMEN EJECUTIVO
    if 'resumen' in sec:
        add_executive_slide(prs, exec_summary, month, year)

    # 2. ÍNDICE DINÁMICO
    idx_sections = []
    idx_n = 1
    if 'trafico' in sec:
        idx_sections.append((f"{idx_n:02d}", "Tráfico Web")); idx_n += 1
    if 'canales' in sec:
        idx_sections.append((f"{idx_n:02d}", "Canales y Audiencia")); idx_n += 1
    if 'search' in sec:
        idx_sections.append((f"{idx_n:02d}", "Rendimiento en Google")); idx_n += 1
    if 'keywords' in sec:
        idx_sections.append((f"{idx_n:02d}", "Keywords y Páginas")); idx_n += 1
    if 'ads' in sec and ads:
        idx_sections.append((f"{idx_n:02d}", "Google Ads")); idx_n += 1
    if 'pagespeed' in sec and pagespeed:
        idx_sections.append((f"{idx_n:02d}", "SEO Técnico")); idx_n += 1
    if 'recomendaciones' in sec and (analysis.get('recommendations') or analysis.get('anomalies')):
        idx_sections.append((f"{idx_n:02d}", "Análisis y Recomendaciones")); idx_n += 1
    if 'conclusiones' in sec:
        idx_sections.append((f"{idx_n:02d}", "Conclusiones y Próximos Pasos"))
    sec_nums = {name: num for num, name in idx_sections}
    add_index_slide(prs, idx_sections)

    # TRÁFICO
    if 'trafico' in sec:
        s_num = sec_nums.get("Tráfico Web", "01")
        s_label = f"{int(s_num)}. Tráfico Web"

        mom_metrics = mom.get('metrics', [])
        if gcharts and mom_metrics:
            from utils import MESES as _MESES
            pm = month - 1 if month > 1 else 12
            img = gcharts.mom_grouped_bars(
                mom_metrics,
                _MESES.get(month, "Actual").capitalize(),
                _MESES.get(pm, "Anterior").capitalize())
            if img:
                add_image_chart_slide(prs, "Comparativa Mes a Mes", img,
                                      mom.get('narrative', ''), s_label, s_num)

        if gcharts and len(historical) >= 3:
            img_trend = gcharts.trend_line_6m(historical)
            if img_trend:
                add_image_chart_slide(prs, "Tendencia — Últimos 6 Meses", img_trend,
                                      "Evolución de usuarios e impresiones en los últimos meses.",
                                      s_label, s_num)

        add_kpi_slide(prs, "Tráfico Web", traffic['narrative'], [
            {"label": "Usuarios",       "value": f"{traffic['users']:,}",               "delta": traffic.get('d_users', '')},
            {"label": "Sesiones",       "value": f"{traffic['sessions']:,}",            "delta": traffic.get('d_sessions', '')},
            {"label": "Páginas vistas", "value": f"{traffic.get('pageviews', 0):,}",    "delta": traffic.get('d_pageviews', '')},
            {"label": "Tasa de rebote", "value": f"{traffic['bounce_rate']:.1f}%",      "delta": ""},
        ], s_label, s_num)
        if include_notes and insights.get('traffic'):
            _add_insight_box(prs.slides[-1], insights['traffic'])

        chart_done = False
        if gcharts and daily_users:
            img = gcharts.ga4_users_line(daily_users)
            if img:
                add_image_chart_slide(prs, "Evolución Diaria de Usuarios", img,
                                      "Usuarios únicos por día durante el mes.",
                                      s_label, s_num)
                chart_done = True
        if not chart_done:
            add_chart_slide(prs, "Evolución del Tráfico", {
                'Usuarios':       traffic['users'],
                'Sesiones':       traffic['sessions'],
                'Páginas vistas': traffic.get('pageviews', 0),
            }, s_label, s_num)

    # CANALES Y AUDIENCIA
    if 'canales' in sec:
        s_num = sec_nums.get("Canales y Audiencia", "02")
        s_label = f"{int(s_num)}. Canales y Audiencia"

        ch_list = channels_d.get('channels', [])
        if ch_list:
            img = gcharts.channels_donut(ch_list) if gcharts else None
            if img:
                add_image_chart_slide(prs, "Canales de Adquisición", img,
                                      channels_d.get('narrative', ''), s_label, s_num)
            else:
                add_channels_slide(prs, "Canales de Adquisición de Tráfico",
                                   ch_list, channels_d['narrative'], s_label, s_num)

        dev_dict = devices_d.get('devices', {})
        if gcharts and (dev_dict or top_countries):
            img_dev = gcharts.hbar(
                [{'label': k, 'value': v} for k, v in dev_dict.items()],
                'label', 'value') if dev_dict else None
            img_geo = gcharts.hbar(top_countries, 'country', 'users',
                                   color="#34A853") if top_countries else None
            if img_dev or img_geo:
                add_double_image_slide(prs, "Dispositivos y Países",
                                       img_dev, "Sesiones por dispositivo",
                                       img_geo, "Usuarios por país", s_label, s_num)
        elif dev_dict:
            add_devices_slide(prs, "Distribución por Dispositivos",
                              dev_dict, devices_d['narrative'], s_label, s_num)

    # RENDIMIENTO EN GOOGLE
    if 'search' in sec:
        s_num = sec_nums.get("Rendimiento en Google", "03")
        s_label = f"{int(s_num)}. Rendimiento en Google"

        add_kpi_slide(prs, "Google Search Console", sc['narrative'], [
            {"label": "Impresiones",   "value": f"{sc['impressions']:,}",  "delta": sc.get('d_impressions', '')},
            {"label": "Clics",         "value": f"{sc['clicks']:,}",       "delta": sc.get('d_clicks', '')},
            {"label": "CTR",           "value": f"{sc['ctr']:.2f}%",       "delta": sc.get('d_ctr', '')},
            {"label": "Posición media","value": f"{sc['position']:.1f}",   "delta": ""},
        ], s_label, s_num)
        if include_notes and insights.get('search'):
            _add_insight_box(prs.slides[-1], insights['search'])

        chart_done = False
        if gcharts and daily_sc:
            img = gcharts.gsc_performance_line(daily_sc)
            if img:
                narr = analysis.get('daily_trends', {}).get('narrative', '')
                add_image_chart_slide(prs, "Evolución en Búsqueda — Clics e Impresiones", img,
                                      narr, s_label, s_num)
                chart_done = True
        if not chart_done:
            add_chart_slide(prs, "Rendimiento en Búsqueda", {
                'Impresiones': sc['impressions'],
                'Clics':       sc['clicks'],
            }, s_label, s_num)

    # KEYWORDS Y PÁGINAS
    if 'keywords' in sec:
        s_num = sec_nums.get("Keywords y Páginas", "04")
        s_label = f"{int(s_num)}. Keywords y Páginas"

        kws = kw.get('top_keywords', [])[:9]
        if kws:
            add_table_slide(prs, "Top Keywords",
                            ["Keyword", "Posición", "Impresiones", "Clics", "CTR"],
                            [[k['query'], f"{k['position']}", f"{k['impressions']:,}",
                              f"{k['clicks']}", f"{k.get('ctr', 0):.1f}%"] for k in kws],
                            kw['narrative'], s_label, s_num,
                            col_widths=[4, 1.2, 1.5, 1, 1])
            if include_notes and insights.get('keywords'):
                _add_insight_box(prs.slides[-1], insights['keywords'])

        add_keyword_movements_slide(prs, kw_moves, s_label, s_num)
        add_opportunities_slide(prs, opps, s_label, s_num)

        pages_ga4 = analysis.get('pages', {}).get('top_pages_ga4', [])
        if pages_ga4:
            add_table_slide(prs, "Páginas Más Visitadas",
                            ["Página", "Visitas", "% del total"],
                            [[pg['path'], f"{pg['views']:,}",
                              f"{pg.get('percentage', 0):.1f}%"] for pg in pages_ga4[:9]],
                            "Páginas con más visitas durante el mes (GA4).",
                            s_label, s_num, col_widths=[6, 1.5, 1.5])

    # GOOGLE ADS (opcional)
    if 'ads' in sec and ads:
        s_num = sec_nums.get("Google Ads", "05")
        s_label = f"{int(s_num)}. Google Ads"
        cur = ads['currency']
        show_conv = ads.get('show_conversions', False)
        last_kpi = ({"label": "Conversiones", "value": f"{ads['conversions']:,.0f}", "delta": ads.get('d_conversions', '')}
                    if show_conv else
                    {"label": "CPC medio", "value": f"{ads['avg_cpc']:.2f} {cur}", "delta": ads.get('d_cpc', '')})
        add_kpi_slide(prs, "Google Ads", ads['narrative'], [
            {"label": "Inversión",    "value": ads['cost_label'],               "delta": ""},
            {"label": "Impresiones",  "value": f"{ads['impressions']:,}",       "delta": ads.get('d_impressions', '')},
            {"label": "Clics",        "value": f"{ads['clicks']:,}",            "delta": ads.get('d_clicks', '')},
            last_kpi,
        ], s_label, s_num)
        if include_notes and ads.get('insight'):
            _add_insight_box(prs.slides[-1], ads['insight'])
        if ads.get('campaigns'):
            _name = lambda c: (c['name'][:42] + '…') if len(c['name']) > 42 else c['name']
            if show_conv:
                cols = ["Campaña", "Inversión", "Clics", "Conv.", "CPA"]
                rows = [[_name(c), f"{c['cost']:,.2f} {cur}", f"{c['clicks']:,}", f"{c['conversions']:,.0f}",
                         f"{c['cpa']:.2f} {cur}" if c.get('cpa') else "—"] for c in ads['campaigns'][:8]]
            else:
                cols = ["Campaña", "Inversión", "Impresiones", "Clics", "CPC"]
                rows = [[_name(c), f"{c['cost']:,.2f} {cur}", f"{c['impressions']:,}", f"{c['clicks']:,}",
                         f"{c['cpc']:.2f} {cur}"] for c in ads['campaigns'][:8]]
            add_table_slide(prs, "Campañas Principales", cols, rows,
                            "Campañas con mayor inversión durante el mes.",
                            s_label, s_num, col_widths=[3.7, 1.5, 1.1, 1.1, 1.3])
        if show_conv and ads.get('conversions_by_action'):
            add_table_slide(prs, "Conversiones por Tipo", ["Acción de conversión", "Conversiones"],
                            [[a['name'], f"{a['conversions']:,.0f}"] for a in ads['conversions_by_action'][:8]],
                            f"Coste por conversión del mes: {ads['cpa_label']}.",
                            s_label, s_num, col_widths=[6.5, 2.2])

    # CAPTURAS DE PANELES
    add_panel_screenshot_slides(prs, client_name)

    # SEO TÉCNICO
    if 'pagespeed' in sec and pagespeed:
        s_num = sec_nums.get("SEO Técnico", "05")
        add_pagespeed_slide(prs, pagespeed, f"{int(s_num)}. SEO Técnico", s_num)

    # ANÁLISIS Y RECOMENDACIONES
    _recs = analysis.get('recommendations', [])
    _anom = analysis.get('anomalies', [])
    if 'recomendaciones' in sec and (_recs or _anom):
        s_num = sec_nums.get("Análisis y Recomendaciones", "06")
        add_recommendations_slide(prs, _recs, _anom, f"{int(s_num)}. Recomendaciones", s_num)

    # CONCLUSIONES
    if 'conclusiones' in sec:
        s_num = sec_nums.get("Conclusiones y Próximos Pasos", "05")
        add_text_slide(prs, "Conclusiones y Próximos Pasos", analysis['conclusions'],
                       f"{int(s_num)}. Conclusiones", s_num)

    # GLOSARIO
    if 'glosario' in sec:
        add_glossary_slide(prs)

    # CONTACTO
    add_contact_slide(prs)

    month_name = MESES_UPPER.get(month, 'MES')
    filename = f"INFORME_SEO_{client_name.upper().replace(' ', '_')}_{month_name}_{year}.pptx"
    filepath = os.path.join(get_output_dir(), filename)
    prs.save(filepath)
    return filepath
