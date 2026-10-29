#!/usr/bin/env python3
"""
Word de un post de blog para enviar al cliente (python-docx).

Contiene el texto completo y, al final, la «Ficha SEO». Nada interno: ni cifras de Ads, ni
volúmenes de búsqueda, ni IDs de WordPress, ni el semáforo.

Carpeta: <OUTPUT_DIR>/para_cliente/<CLIENTE>/POST_<AAAA>-<MM>_<slug>_v<N>.docx
"""

import glob
import os
import re
from datetime import date

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import seo_posts as sp
from utils import MESES, get_output_dir

NAVY = RGBColor(0x1E, 0x1E, 0x6E)
MUTE = RGBColor(0x7A, 0x7C, 0x9C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Calibri"
INLINE = re.compile(r"(\[[^\]]+\]\([^)\s]+\)|\*\*[^*]+\*\*|\*[^*\n]+\*)")


# ── Utilidades de formato ──────────────────────────────────────
def _style_font(style, size=None, color=None, bold=None):
    style.font.name = FONT
    rf = style.element.get_or_add_rPr().get_or_add_rFonts()
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rf.attrib.pop(qn(a), None)
    rf.set(qn("w:eastAsia"), FONT)
    if size:
        style.font.size = Pt(size)
    if color is not None:
        style.font.color.rgb = color
    if bold is not None:
        style.font.bold = bold


def _shade(cell, fill: str):
    tc = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    tc.append(shd)


def _hyperlink(par, text: str, url: str):
    r_id = par.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hl = OxmlElement("w:hyperlink")
    hl.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color"); color.set(qn("w:val"), "1E1EBE"); rpr.append(color)
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rpr.append(u)
    run.append(rpr)
    t = OxmlElement("w:t"); t.text = text; t.set(qn("xml:space"), "preserve")
    run.append(t); hl.append(run)
    par._p.append(hl)


def _inline(par, text: str, base: str = ""):
    """Texto con enlaces [a](u), **negrita** y *cursiva*."""
    for part in INLINE.split(text or ""):
        if not part:
            continue
        m = re.fullmatch(r"\[([^\]]+)\]\(([^)\s]+)\)", part)
        if m:
            _hyperlink(par, m.group(1), _abs(m.group(2), base))
        elif part.startswith("**") and part.endswith("**"):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            par.add_run(part[1:-1]).italic = True
        else:
            par.add_run(part)


def _abs(url: str, base: str) -> str:
    return base.rstrip("/") + url if url.startswith("/") and base else url


def _box(doc, fill: str, lines: list, base: str = ""):
    """Recuadro de una celda con varias líneas [(texto, negrita)]."""
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    cell = t.rows[0].cells[0]
    _shade(cell, fill)
    first = True
    for text, bold in lines:
        par = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        if bold:
            par.add_run(text).bold = True
        else:
            _inline(par, text, base)
    doc.add_paragraph()


# ── Documento ──────────────────────────────────────────────────
def build_docx(post: dict, cliente: str, site_base: str, version: int, mes: int, anio: int, path: str):
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2.3)
    sec.top_margin = sec.bottom_margin = Cm(2.2)
    _style_font(doc.styles["Normal"], 11)
    _style_font(doc.styles["Heading 1"], 24, NAVY, True)
    _style_font(doc.styles["Heading 2"], 16, NAVY, True)
    _style_font(doc.styles["Heading 3"], 13, NAVY, True)
    _style_font(doc.styles["Heading 4"], 11.5, NAVY, True)
    doc.styles["Normal"].paragraph_format.space_after = Pt(8)

    hp = sec.header.paragraphs[0]
    hr = hp.add_run("Inpulsia · Propuesta de contenido")
    hr.font.size = Pt(9); hr.font.color.rgb = MUTE
    fp = sec.footer.paragraphs[0]
    fr = fp.add_run(f"Borrador para revisión · versión {version} · {MESES.get(mes, mes).capitalize()} {anio}")
    fr.font.size = Pt(9); fr.font.color.rgb = MUTE

    label = doc.add_paragraph()
    lr = label.add_run(f"PROPUESTA DE CONTENIDO · {cliente.replace('_', ' ').upper()} · BORRADOR V{version}")
    lr.font.size = Pt(9); lr.font.color.rgb = MUTE; lr.bold = True
    doc.add_heading(post.get("titulo", ""), level=1)
    if post.get("extracto"):
        ex = doc.add_paragraph()
        er = ex.add_run(post["extracto"]); er.italic = True; er.font.color.rgb = MUTE

    for b in sp.parse_content(post.get("contenido", "")):
        t = b["t"]
        if t == "p":
            _inline(doc.add_paragraph(), b["text"], site_base)
        elif t in ("h2", "h3", "h4", "h5", "h6"):
            doc.add_heading(b["text"], level=min(int(t[1]) , 4))
        elif t == "h1":
            doc.add_heading(b["text"], level=1)
        elif t == "ul":
            for it in b["items"]:
                _inline(doc.add_paragraph(style="List Bullet"), it, site_base)
        elif t == "ol":
            for n, it in enumerate(b["items"], 1):
                par = doc.add_paragraph()
                par.paragraph_format.left_indent = Cm(0.9)
                par.paragraph_format.first_line_indent = Cm(-0.6)
                par.paragraph_format.space_after = Pt(4)
                par.add_run(f"{n}. ")
                _inline(par, it, site_base)
        elif t == "table" and b["rows"]:
            cols = max(len(r) for r in b["rows"])
            tbl = doc.add_table(rows=len(b["rows"]), cols=cols)
            tbl.style = "Table Grid"
            for ri, row in enumerate(b["rows"]):
                for ci in range(cols):
                    cell = tbl.rows[ri].cells[ci]
                    txt = row[ci] if ci < len(row) else ""
                    par = cell.paragraphs[0]
                    if ri == 0:
                        _shade(cell, "1E1E6E")
                        run = par.add_run(sp.plain(txt)); run.bold = True; run.font.color.rgb = WHITE
                    else:
                        _inline(par, txt, site_base)
            doc.add_paragraph()
        elif t == "img":
            lines = [("Imagen sugerida", True), (b["ref"] or "(a elegir)", False)]
            if b["alt"]:
                lines.append((f"Texto alternativo: {b['alt']}", False))
            _box(doc, "F1F2FA", lines)
        elif t == "quote":
            _box(doc, "F4FFD6", [(b["text"], False)], site_base)

    cta = doc.add_paragraph()
    cta.add_run("¿Cómo revisar este documento?").bold = True
    doc.add_paragraph("Indícanos qué cambiarías del texto (puedes usar comentarios o control de cambios). "
                      "Cuando lo apruebes, lo publicaremos en la web. La ficha SEO de la página siguiente es informativa: "
                      "la aplicamos nosotros.")

    # Ficha SEO
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading("Ficha SEO", level=2)
    tseo, meta = post.get("titulo_seo", ""), post.get("meta_description", "")
    internal, external = sp.split_links(sp.links_in(sp.parse_content(post.get("contenido", ""))), site_base)
    slug = (post.get("slug") or "").strip("/")
    rows = [
        ("Título SEO", f"{tseo} ({len(tseo)} caracteres)" if tseo else ""),
        ("Meta descripción", f"{meta} ({len(meta)} caracteres)" if meta else ""),
        ("Dirección (URL)", f"{site_base.rstrip('/')}/{slug}/" if slug and site_base else (f"/{slug}/" if slug else "")),
        ("Palabra clave principal", post.get("keyword_principal", "")),
        ("Palabras clave secundarias", ", ".join(post.get("keywords_secundarias") or [])),
        ("Categoría", post.get("categoria", "")),
        ("Etiquetas", ", ".join(post.get("etiquetas") or [])),
        ("Tipo de schema", ", ".join(post.get("schema") or [])),
        ("Imagen destacada (texto alternativo)", post.get("imagen_destacada_alt", "")),
        ("Enlaces internos", "\n".join(f"{a} → {_abs(u, site_base)}" for a, u in internal)),
        ("Enlaces externos", "\n".join(f"{a} → {u}" for a, u in external)),
    ]
    rows = [r for r in rows if r[1]]
    tbl = doc.add_table(rows=len(rows), cols=2)
    tbl.style = "Table Grid"
    for i, (k, v) in enumerate(rows):
        _shade(tbl.rows[i].cells[0], "F1F2FA")
        tbl.rows[i].cells[0].paragraphs[0].add_run(k).bold = True
        tbl.rows[i].cells[1].paragraphs[0].add_run(v)
        tbl.rows[i].cells[0].width = Cm(5.2)
        tbl.rows[i].cells[1].width = Cm(11.2)

    doc.core_properties.title = post.get("titulo", "")
    doc.core_properties.author = "Inpulsia"
    doc.save(path)


def _safe_slug(s: str) -> str:
    return re.sub(r"[^a-z0-9-]+", "-", (s or "post").lower()).strip("-") or "post"


def create_word(post: dict, cliente: str, site_base: str = "", mes: int = None, anio: int = None) -> dict:
    """Guarda el Word y devuelve {'path', 'version'}. Cada llamada crea una versión nueva (v1, v2…)."""
    today = date.today()
    mes, anio = int(mes or today.month), int(anio or today.year)
    folder = os.path.join(get_output_dir(), "para_cliente", re.sub(r"[^A-Za-z0-9_-]", "", cliente.upper()))
    os.makedirs(folder, exist_ok=True)
    slug = _safe_slug(post.get("slug") or post.get("titulo"))
    prefix = f"POST_{anio}-{mes:02d}_{slug}_v"
    versions = [int(m.group(1)) for f in glob.glob(os.path.join(folder, prefix + "*.docx"))
                if (m := re.search(r"_v(\d+)\.docx$", f))]
    version = max(versions or [0]) + 1
    path = os.path.join(folder, f"{prefix}{version}.docx")
    build_docx(post, cliente, site_base, version, mes, anio, path)
    return {"path": path, "version": version}
